"""선행연구 구성 분석 (IMRaD 과학 논문: 서론의 선행연구 토픽 + 결과·논의의 선행연구 활용).

  build    : corpus/*.json(⟦c:…⟧ 인용 표지 포함) -> spaCy 문장 분할·동사 추출 -> 엑셀
             - 서론_인용문장: 인용 문장마다 문단, 펼친 문헌 번호, 통합 인용 여부, 자동 토픽(=문단)
             - 결과논의_인용문장: 인용 문장 + 앞·뒤 문장, 자동 기능 판정
  dump     : Claude 검토용 텍스트 (서론은 논문별 문단 전체, 결과·논의는 인용 문장)
  stats    : 판정(연구자 > Claude > 자동)을 읽어 논문별·분야 요약 시트 재계산
  baseline : 진단 아티팩트용 기준선 JSON

사용:
  python litreview.py build    <corpus_dir> <out.xlsx> [review.json]
  python litreview.py dump     <out.xlsx> <corpus_dir> <candidates.txt>
  python litreview.py stats    <out.xlsx>
  python litreview.py baseline <out.xlsx> <baseline.json>
"""
import sys, os, re, json, glob, statistics
from collections import Counter, defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SECS = ['INTRO', 'M', 'RD', 'CONCL']
SEC_NAME = {'INTRO': '서론', 'M': '방법', 'RD': '결과·논의', 'CONCL': '결론'}
FUNCS = ['비교·일치', '비교·차이', '해석 근거', '방법·자료 출처', '배경 정보']
ORGS = ['주제 통합', '논문별 나열', '혼합']
INTRO_SHEET, TOPIC_SHEET, RD_SHEET, SIZE_SHEET = '서론_인용문장', '서론_토픽', '결과논의_인용문장', '섹션크기'

CIT = re.compile(r'⟦c:([^⟧]*)⟧')
INTEGRAL = re.compile(r"\b[A-Z][a-zA-Z\-]+(?:\s+(?:and|&)\s+[A-Z][a-zA-Z\-]+)?\s+(?:et al\.?|and co-?workers|and colleagues)"
                      r"|\b(?:the\s+)?groups?\s+of\s+[A-Z]|\b[A-Z][a-z]+(?:'s|’s)\s+group|\b[A-Z][a-z]+\s+\(\d{4}[a-z]?\)")

FUNC_RULES = [
    ('비교·차이', r"in contrast|\bunlike\b|contrary to|\bdiffer|in disagreement|inconsistent with|whereas[^.]*report"),
    ('비교·일치', r"consistent with|in (?:good |excellent |reasonable )?agreement|agrees? (?:well )?with|in line with|in accordance with|"
                  r"similar to (?:that|those|the one)s?[^.]*(?:report|observ|describ)|as (?:previously |already |was )?(?:reported|described|observed|found|shown)|"
                  r"confirms? (?:the|earlier|previous)|corroborat|comparable (?:to|with)|in keeping with|not unusual"),
    ('방법·자료 출처', r"according to|following (?:the |a )?(?:literature|reported|published|previous|known)|(?:reported|literature|published|known|standard) "
                     r"(?:procedure|method|protocol|condition)s?|as described (?:previously|elsewhere|in)|were (?:taken|adopted|obtained) from|"
                     r"data (?:from|taken)|(?:method|approach) of|level of theory|(?:at|using) the \S+ level|basis set|functional|Gaussian|program|software"),
    ('해석 근거', r"(?:is|are|has been|have been|was|were) (?:well[- ])?(?:known|established|shown|demonstrated|documented|reported) (?:to|that)|"
                 r"it is (?:well[- ])?known|attributed to|ascribed to|explained by|rationaliz|due to"),
]


# ---------------------------------------------------------------- 인용 표지 처리
def expand(body):
    """'1−4' -> {1,2,3,4}; '3a,b' -> {3}; '7a,9' -> {7,9}"""
    refs = set(); last = None
    for part in re.split(r'\s*,\s*', body):
        m = re.match(r'^(\d*)([a-h]?)\s*(?:[−–-]\s*(\d*)([a-h]?))?$', part.strip())
        if not m: continue
        a, la, b, lb = m.groups()
        if a: last = int(a)
        if a and b:
            refs.update(range(int(a), int(b) + 1))
        elif last is not None:
            refs.add(last)
    refs.discard(0)
    return refs


def strip_markers(p):
    """문단 -> (표지 없는 본문, [(본문 내 위치, 문헌 집합, 원래 표지)])"""
    out = []; marks = []; pos = 0; clen = 0
    for m in CIT.finditer(p):
        seg = p[pos:m.start()]; out.append(seg); clen += len(seg)
        marks.append((clen, expand(m.group(1)), m.group(1)))
        pos = m.end()
    out.append(p[pos:])
    return ''.join(out), marks


def auto_function(sent):
    s = sent.lower()
    for name, pat in FUNC_RULES:
        if re.search(pat, s, re.I): return name, f'단서: {re.search(pat, s, re.I).group(0)[:30]}'
    return '배경 정보', '단서 없음 — 일반 배경 진술로 추정'


def lex_verbs(span):
    return [t.lemma_.lower() for t in span if t.pos_ == 'VERB' and t.lemma_.isalpha() and len(t.lemma_) > 1]


def is_word(t):
    return not (t.is_punct or t.is_space) and any(ch.isalnum() for ch in t.text)


# ---------------------------------------------------------------- build
def build(corpus_dir, out_xlsx, review_path=None):
    import spacy
    nlp = spacy.load('en_core_web_sm')
    docs = [json.load(open(f, encoding='utf-8')) for f in sorted(glob.glob(os.path.join(corpus_dir, '*.json')))]
    docs.sort(key=lambda d: (d['year'] or 0, d['id']))
    intro_rows, rd_rows, sizes = [], [], []
    for i, d in enumerate(docs):
        pid = f'P{i + 1:02d}'
        for sec in SECS:
            paras = d['sections'].get(sec, [])
            words = 0; refs_all = set(); n_marks = 0; all_sents = []
            for pi, p in enumerate(paras, 1):
                text, marks = strip_markers(p)
                doc = nlp(text)
                words += sum(1 for t in doc if is_word(t))
                sents = list(doc.sents)
                for si, s in enumerate(sents):
                    # 문장 끝 마침표 뒤에 붙은 위첨자도 그 문장 것으로 본다
                    end = sents[si + 1].start_char if si + 1 < len(sents) else len(text) + 1
                    ms = [m for m in marks if s.start_char < m[0] <= end]
                    all_sents.append((pi, s, ms))
            for k, (pi, s, ms) in enumerate(all_sents):
                n_marks += len(ms)
                refs = set().union(*[m[1] for m in ms]) if ms else set()
                refs_all |= refs
                if not ms: continue
                row = dict(pid=pid, file=d['id'], sec=sec, para=pi, sent_no=k + 1, marks=len(ms),
                           bodies=', '.join(m[2] for m in ms), refs=sorted(refs), sent=s.text.strip(),
                           integral=bool(INTEGRAL.search(s.text)), verbs=lex_verbs(s))
                if sec == 'INTRO':
                    intro_rows.append(row)
                elif sec == 'RD':
                    row['prev'] = all_sents[k - 1][1].text.strip() if k > 0 else ''
                    row['next'] = all_sents[k + 1][1].text.strip() if k + 1 < len(all_sents) else ''
                    row['auto'], row['why'] = auto_function(s.text)
                    rd_rows.append(row)
            sizes.append(dict(pid=pid, file=d['id'], sec=sec, words=words, paras=len(paras), marks=n_marks,
                              refs=sorted(refs_all), n_refs=d.get('n_refs'),
                              cite_paras=len({pi for pi, s, ms in all_sents if ms})))
        print(pid, d['id'], 'refs', d.get('n_refs'), {s['sec']: (s['words'], s['marks'], len(s['refs'])) for s in sizes if s['pid'] == pid})
    write_workbook(out_xlsx, intro_rows, rd_rows, sizes, review_path)
    stats(out_xlsx)


HDR_FONT = Font(bold=True, color='FFFFFF')
HDR_FILL = PatternFill('solid', fgColor='305496')


def header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i); cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'


def write_workbook(path, intro_rows, rd_rows, sizes, review_path):
    review = json.load(open(review_path, encoding='utf-8')) if review_path and os.path.exists(review_path) else {}
    rv_intro, rv_rd, rv_org = review.get('intro', {}), review.get('rd', {}), review.get('topic_org', {})
    wb = Workbook()
    ws = wb.active; ws.title = '안내'
    for line in README: ws.append([line])
    ws.column_dimensions['A'].width = 140

    ws = wb.create_sheet(INTRO_SHEET)
    header(ws, ['ID', '논문', '문단', '문장번호', '인용 표지', '인용 문헌(펼침)', '문헌 수', '통합 인용', '인용 문장', '문장의 동사',
                '자동 토픽', 'Claude 토픽', 'Claude 메모', '연구자 토픽', '메모'],
           [5, 6, 5, 8, 12, 18, 7, 8, 80, 28, 9, 26, 26, 18, 20])
    for i, r in enumerate(intro_rows, 1):
        topic, memo = rv_intro.get(str(i), ['', ''])
        ws.append([i, r['pid'], r['para'], r['sent_no'], r['bodies'], ', '.join(map(str, r['refs'])), len(r['refs']),
                   'Y' if r['integral'] else '', r['sent'], ', '.join(r['verbs']), f"¶{r['para']}", topic, memo, '', ''])
    for row in ws.iter_rows(min_row=2): row[8].alignment = Alignment(wrap_text=True, vertical='top')
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet(TOPIC_SHEET)       # stats가 토픽 목록을 다시 만들고, 조직 유형 판정은 보존
    header(ws, ['논문', '토픽', '문헌 수', '인용 문장 수', '문헌 번호', 'Claude 조직 유형', '연구자 조직 유형'], [6, 30, 8, 10, 30, 14, 14])
    ws.append(['_org_review', json.dumps(rv_org, ensure_ascii=False)])   # Claude 조직 유형 원본 (stats에서 사용)

    ws = wb.create_sheet(RD_SHEET)
    header(ws, ['ID', '논문', '문장번호', '인용 표지', '인용 문헌(펼침)', '통합 인용', '앞 문장', '인용 문장', '뒤 문장', '문장의 동사',
                '자동 기능', '판정 근거', 'Claude 기능', 'Claude 메모', '연구자 기능', '메모'],
           [5, 6, 8, 12, 16, 8, 40, 70, 40, 26, 12, 22, 12, 26, 14, 20])
    for i, r in enumerate(rd_rows, 1):
        fn, memo = rv_rd.get(str(i), ['', ''])
        ws.append([i, r['pid'], r['sent_no'], r['bodies'], ', '.join(map(str, r['refs'])), 'Y' if r['integral'] else '',
                   r['prev'], r['sent'], r['next'], ', '.join(r['verbs']), r['auto'], r['why'], fn, memo, '', ''])
    for row in ws.iter_rows(min_row=2):
        for c in row[6:9]: c.alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"' + ','.join(FUNCS) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'O2:O{max(2, len(rd_rows) + 1)}')
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet(SIZE_SHEET)
    header(ws, ['ID', '논문', '섹션', '단어 수', '문단 수', '인용 문단 수', '인용 표지 수', '고유 문헌 수', '문헌 번호', '전체 참고문헌'],
           [6, 12, 10, 9, 8, 10, 10, 10, 40, 10])
    for s in sizes:
        ws.append([s['pid'], s['file'], SEC_NAME[s['sec']], s['words'], s['paras'], s['cite_paras'], s['marks'], len(s['refs']),
                   ', '.join(map(str, s['refs'])), s['n_refs']])
    wb.save(path)


# ---------------------------------------------------------------- stats
def pick(*vals):
    for v in vals:
        if isinstance(v, str) and v.strip(): return v.strip()
    return None


def summarize(vals):
    v = [x for x in vals if x is not None]
    if not v: return dict(mean=None, sd=None, median=None, lo=None, hi=None, k=0, n=0)
    return dict(mean=statistics.mean(v), sd=statistics.stdev(v) if len(v) > 1 else 0.0, median=statistics.median(v),
                lo=min(v), hi=max(v), k=sum(1 for x in v if x > 0), n=len(v))


def load_tables(wb):
    size = {}
    for r in wb[SIZE_SHEET].iter_rows(min_row=2, values_only=True):
        if r[0]: size[(r[0], r[2])] = dict(file=r[1], words=r[3], paras=r[4], cite_paras=r[5], marks=r[6], nrefs_sec=r[7],
                                         refs=set(int(x) for x in str(r[8] or '').split(', ') if x), n_refs=r[9])
    intro = []
    for r in wb[INTRO_SHEET].iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        intro.append(dict(id=r[0], pid=r[1], para=r[2], marks=1 if r[4] else 0, nm=len(str(r[4]).split(', ')) if r[4] else 0,
                          refs=[int(x) for x in str(r[5] or '').split(', ') if x], integral=r[7] == 'Y', sent=r[8],
                          verbs=[v for v in str(r[9] or '').split(', ') if v], topic=pick(r[13], r[11], r[10])))
    rd = []
    for r in wb[RD_SHEET].iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        fn = pick(r[14], r[12], r[10])
        rd.append(dict(id=r[0], pid=r[1], nm=len(str(r[3]).split(', ')) if r[3] else 0,
                       refs=[int(x) for x in str(r[4] or '').split(', ') if x], integral=r[5] == 'Y',
                       verbs=[v for v in str(r[9] or '').split(', ') if v], func=fn if fn in FUNCS else '배경 정보',
                       src='연구자' if pick(r[14]) else 'Claude' if pick(r[12]) else '자동'))
    return size, intro, rd


def compute(wb):
    size, intro, rd = load_tables(wb)
    papers = sorted({p for p, _ in size})
    ws = wb[TOPIC_SHEET]
    org_claude = {}; org_user = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] == '_org_review': org_claude = json.loads(r[1] or '{}')
        elif r[0] and r[1] and pick(r[6]): org_user[f'{r[0]}|{r[1]}'] = r[6].strip()
    per = {}
    for p in papers:
        I = [x for x in intro if x['pid'] == p]; R = [x for x in rd if x['pid'] == p]
        topics = defaultdict(lambda: dict(refs=set(), sents=0, paras=set()))
        for x in I:
            if x['topic'].startswith('(제외)'): continue       # 본 연구 서술(M3) 속 인용 등은 선행연구 토픽이 아님
            t = topics[x['topic']]; t['refs'] |= set(x['refs']); t['sents'] += 1; t['paras'].add(x['para'])
        tlist = []
        for name, t in topics.items():
            org = org_user.get(f'{p}|{name}') or org_claude.get(p, {}).get(name)
            tlist.append(dict(name=name, refs=sorted(t['refs']), n=len(t['refs']), sents=t['sents'], paras=sorted(t['paras']), org=org))
        tlist.sort(key=lambda t: t['paras'][0])
        si, sr = size[(p, '서론')], size[(p, '결과·논의')]
        nrefs = si['n_refs'] or None
        intro_refs, rd_refs = si['refs'], sr['refs']
        funcs = Counter(x['func'] for x in R)
        per[p] = dict(
            file=si['file'], n_refs=nrefs,
            intro=dict(words=si['words'], paras=si['paras'], cite_paras=si['cite_paras'], marks=si['marks'], refs=len(intro_refs),
                       density=si['marks'] / si['words'] * 100 if si['words'] else 0,
                       topics=tlist, n_topics=len(tlist), refs_per_topic=statistics.mean([t['n'] for t in tlist]) if tlist else 0,
                       integral=sum(x['integral'] for x in I) / len(I) * 100 if I else 0, cite_sents=len(I),
                       verbs=Counter(v for x in I for v in x['verbs']),
                       org=Counter(t['org'] for t in tlist if t['org'])),
            rd=dict(words=sr['words'], marks=sr['marks'], refs=len(rd_refs), cite_sents=len(R),
                    density=sr['marks'] / sr['words'] * 1000 if sr['words'] else 0,
                    new_refs=len(rd_refs - intro_refs), shared_refs=len(rd_refs & intro_refs),
                    funcs=funcs, func_pct={f: funcs[f] / len(R) * 100 if R else 0 for f in FUNCS},
                    integral=sum(x['integral'] for x in R) / len(R) * 100 if R else 0,
                    verbs=Counter(v for x in R for v in x['verbs'])),
            share={s: (len(size[(p, s)]['refs']) / nrefs * 100 if nrefs else None) for s in ['서론', '방법', '결과·논의', '결론']},
        )
    return papers, per, intro, rd


def stats(path):
    wb = load_workbook(path)
    papers, per, intro, rd = compute(wb)
    for name in ['논문별 요약', '분야 요약', '인용 문장 동사', '결과논의 기능']:
        if name in wb.sheetnames: del wb[name]
    # 서론_토픽 시트 다시 쓰기 (연구자 조직 유형 열은 보존)
    ws = wb[TOPIC_SHEET]
    keep = {(r[0], r[1]): r[6] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[0] != '_org_review'}
    org_row = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == '_org_review']
    ws.delete_rows(2, ws.max_row)
    for p in papers:
        for t in per[p]['intro']['topics']:
            ws.append([p, t['name'], t['n'], t['sents'], ', '.join(map(str, t['refs'])), t['org'] or '', keep.get((p, t['name'])) or ''])
    ws.append(list(org_row[0][:2]) if org_row else ['_org_review', '{}'])
    dv = DataValidation(type='list', formula1='"' + ','.join(ORGS) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'G2:G{ws.max_row}')

    ws = wb.create_sheet('논문별 요약')
    cols = ['ID', '논문', '전체 참고문헌', '서론 단어', '서론 문단', '서론 인용 표지', '서론 고유 문헌', '인용 밀도(/100단어)', '토픽 수',
            '토픽당 문헌(평균)', '토픽별 문헌 수', '통합 인용(%)', '토픽 조직', '결과·논의 단어', '결과·논의 인용 표지', '결과·논의 고유 문헌',
            '결과·논의 인용 밀도(/1,000단어)', '결과·논의에서 새로 인용', '서론과 공유'] + [f'기능:{f}(%)' for f in FUNCS] + \
           ['참고문헌 중 서론 인용(%)', '참고문헌 중 결과·논의 인용(%)', '참고문헌 중 방법 인용(%)']
    header(ws, cols, [5, 11] + [10] * (len(cols) - 2))
    for p in papers:
        a, b = per[p]['intro'], per[p]['rd']
        ws.append([p, per[p]['file'], per[p]['n_refs'], a['words'], a['paras'], a['marks'], a['refs'], rnd(a['density']), a['n_topics'],
                   rnd(a['refs_per_topic']), ' / '.join(str(t['n']) for t in a['topics']), rnd(a['integral'], 0),
                   ', '.join(f'{k} {v}' for k, v in a['org'].items()), b['words'], b['marks'], b['refs'], rnd(b['density']), b['new_refs'],
                   b['shared_refs']] + [rnd(b['func_pct'][f], 0) for f in FUNCS] +
                  [rnd(per[p]['share']['서론'], 0), rnd(per[p]['share']['결과·논의'], 0), rnd(per[p]['share']['방법'], 0)])

    ws = wb.create_sheet('분야 요약')
    header(ws, ['영역', '지표', '평균', '표준편차', '중앙값', '최소', '최대', '범위', '출현 논문 수'], [10, 30, 9, 9, 9, 9, 9, 16, 12])
    for area, key, label, get in METRICS:
        s = summarize([get(per[p]) for p in papers])
        ws.append([area, label, rnd(s['mean']), rnd(s['sd']), rnd(s['median']), rnd(s['lo']), rnd(s['hi']),
                   f"{fmt(s['lo'])}–{fmt(s['hi'])}", f"{s['k']}/{s['n']}"])

    ws = wb.create_sheet('인용 문장 동사')
    header(ws, ['영역', '동사', '토큰', '출현 논문 수', '논문별'], [10, 14, 8, 12, 50])
    for area, k in [('서론', 'intro'), ('결과·논의', 'rd')]:
        tot = Counter(); kk = Counter(); byp = defaultdict(dict)
        for p in papers:
            for v, n in per[p][k]['verbs'].items(): tot[v] += n; kk[v] += 1; byp[v][p] = n
        for v, n in sorted(tot.items(), key=lambda x: (-kk[x[0]], -x[1]))[:25]:
            ws.append([area, v, n, f'{kk[v]}/{len(papers)}', ', '.join(f'{p}:{c}' for p, c in byp[v].items())])

    ws = wb.create_sheet('결과논의 기능')
    header(ws, ['기능', '문장 수', '출현 논문 수'] + papers, [16, 8, 12] + [7] * len(papers))
    for f in FUNCS:
        row = [per[p]['rd']['funcs'][f] for p in papers]
        ws.append([f, sum(row), f'{sum(1 for x in row if x)}/{len(papers)}'] + row)
    src = Counter(x['src'] for x in rd)
    ws.append([]); ws.append([f'판정 출처(결과·논의 기능): {dict(src)} / 서론 토픽은 연구자 토픽 > Claude 토픽 > 자동(문단) 순'])
    wb.save(path)
    print('통계 갱신:', path, '| 토픽 수', {p: per[p]['intro']['n_topics'] for p in papers}, '| 기능 판정 출처', dict(src))


METRICS = [
    ('서론', 'n_topics', '선행연구 토픽 수', lambda q: q['intro']['n_topics']),
    ('서론', 'refs_per_topic', '토픽당 문헌 수(평균)', lambda q: q['intro']['refs_per_topic']),
    ('서론', 'refs', '서론 고유 인용 문헌 수', lambda q: q['intro']['refs']),
    ('서론', 'marks', '서론 인용 표지 수', lambda q: q['intro']['marks']),
    ('서론', 'density', '인용 밀도(/100단어)', lambda q: q['intro']['density']),
    ('서론', 'integral', '통합 인용 문장 비율(%)', lambda q: q['intro']['integral']),
    ('결과·논의', 'refs', '결과·논의 고유 인용 문헌 수', lambda q: q['rd']['refs']),
    ('결과·논의', 'marks', '결과·논의 인용 표지 수', lambda q: q['rd']['marks']),
    ('결과·논의', 'density', '인용 밀도(/1,000단어)', lambda q: q['rd']['density']),
    ('결과·논의', 'new_refs', '서론에 없던 새 문헌 수', lambda q: q['rd']['new_refs']),
] + [('결과·논의', 'func:' + f, f'기능: {f}(%)', (lambda f: lambda q: q['rd']['func_pct'][f])(f)) for f in FUNCS] + [
    ('참고문헌', 'share_intro', '서론에서 인용된 비율(%)', lambda q: q['share']['서론']),
    ('참고문헌', 'share_rd', '결과·논의에서 인용된 비율(%)', lambda q: q['share']['결과·논의']),
]


# ---------------------------------------------------------------- dump / baseline
def dump(path, corpus_dir, out_txt):
    wb = load_workbook(path)
    _, intro, rd = load_tables(wb)
    docs = [json.load(open(f, encoding='utf-8')) for f in sorted(glob.glob(os.path.join(corpus_dir, '*.json')))]
    docs.sort(key=lambda d: (d['year'] or 0, d['id']))
    with open(out_txt, 'w', encoding='utf-8') as f:
        f.write('=== 서론: 문단 전체(⟦c:…⟧ = 인용 표지) + 인용 문장 ID ===\n')
        for i, d in enumerate(docs):
            pid = f'P{i + 1:02d}'
            f.write(f'\n## {pid} {d["id"]}\n')
            for pi, p in enumerate(d['sections'].get('INTRO', []), 1):
                f.write(f'¶{pi}: {p}\n')
            for x in [x for x in intro if x['pid'] == pid]:
                f.write(f'  #{x["id"]} ¶{x["para"]} 문헌 {x["refs"]} | {x["sent"][:160]}\n')
        f.write('\n=== 결과·논의 인용 문장 ===\n')
        ws = wb[RD_SHEET]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None: continue
            f.write(f'#{r[0]} {r[1]} 문헌 {r[4]} [자동 {r[10]}: {r[11]}]\n  앞: {r[6]}\n  ▶ {r[7]}\n  뒤: {r[8]}\n\n')
    print('→', out_txt)


def baseline(path, out_json):
    wb = load_workbook(path)
    papers, per, intro, rd = compute(wb)
    out = dict(papers=[dict(pid=p, file=per[p]['file'], n_refs=per[p]['n_refs']) for p in papers], metrics={}, topics={}, verbs={}, funcs=FUNCS)
    for area, key, label, get in METRICS:
        out['metrics'][f'{area}:{key}'] = dict(label=label, values=[rnd(get(per[p]), 3) for p in papers])
    for p in papers:
        out['topics'][p] = [dict(name=t['name'], n=t['n'], org=t['org']) for t in per[p]['intro']['topics']]
    for k in ['intro', 'rd']:
        tot = Counter(); kk = Counter()
        for p in papers:
            for v, n in per[p][k]['verbs'].items(): tot[v] += n; kk[v] += 1
        out['verbs'][k] = [[v, kk[v], tot[v]] for v, _ in sorted(tot.items(), key=lambda x: (-kk[x[0]], -x[1]))[:40]]
    json.dump(out, open(out_json, 'w', encoding='utf-8'), ensure_ascii=False)
    print('→', out_json)


def rnd(x, d=2):
    return None if x is None else round(x, d)


def fmt(x):
    return '—' if x is None else f'{x:.2f}'


README = [
    '선행연구 구성 분석 — 작업 안내 (IMRaD 과학 논문용)',
    '',
    '과학 논문은 선행연구 절이 따로 없고 서론에 녹아 있습니다. 그래서 두 곳을 봅니다.',
    '  서론: 인용 문장을 선행연구 토픽으로 묶어 토픽 수, 토픽당 문헌 수, 인용 문장의 주요 동사, 토픽 조직(주제 통합/논문별 나열/혼합)을 셉니다.',
    '  결과·논의: 인용 문장이 선행연구를 어디에 쓰는지(비교·일치, 비교·차이, 해석 근거, 방법·자료 출처, 배경 정보)와 동사, 서론에 없던 새 문헌 수를 셉니다.',
    '인용: 번호식 위첨자(⟦c:1−4⟧)를 펼쳐 셉니다. 1−4는 4편, 3a,b는 1편. 참고문헌 최대 번호보다 큰 위첨자는 인용이 아닙니다.',
    '통합 인용: 저자가 문장 성분으로 나오는 인용(Kinjo and co-workers, the groups of Bertrand) — 정규식 자동 표시.',
    '',
    '[판정 우선순위] 연구자 > Claude > 자동',
    "  서론_인용문장: N열 '연구자 토픽'에 토픽 이름을 적으면 그 토픽으로 묶입니다(같은 이름 = 같은 토픽). 비우면 Claude 토픽, 그것도 없으면 문단(¶n).",
    "  서론_토픽: G열 '연구자 조직 유형'(주제 통합/논문별 나열/혼합).",
    "  결과논의_인용문장: O열 '연구자 기능'(드롭다운).",
    '  저장 후 다음을 실행하면 요약 시트가 다시 계산됩니다:  <venv python> <스킬 폴더>/scripts/litreview.py stats <이 엑셀>',
]

if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'build': build(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else None)
    elif cmd == 'dump': dump(sys.argv[2], sys.argv[3], sys.argv[4])
    elif cmd == 'stats': stats(sys.argv[2])
    elif cmd == 'baseline': baseline(sys.argv[2], sys.argv[3])
