"""Shared analysis for manuscripts, PDF jobs and field reference corpora."""
from collections import Counter
from functools import lru_cache
import importlib.util
import json
from pathlib import Path
import re
import statistics
import tempfile

ROOT = Path(__file__).resolve().parents[1]
NAMES = {'INTRO': '서론', 'M': '방법', 'RD': '결과·논의', 'CONCL': '결론'}


@lru_cache
def script(name):
    group = 'litreview-structure' if name in ('litreview', 'extract_citations') else 'hedge-verb-sections'
    spec = importlib.util.spec_from_file_location(name, ROOT / 'GNU_AI_Team' / group / 'scripts' / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@lru_cache
def language():
    import spacy
    return spacy.load('en_core_web_sm')


def paragraphs(text):
    return [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]


def markers(text):
    return re.sub(r'\[(\d+[a-h]?(?:\s*[,−–-]\s*(?:\d+[a-h]?|[a-h]))*)\]', r'⟦c:\1⟧', text)


def split_sections(text):
    sections = {k: '' for k in NAMES}
    current = None
    front = []
    for line in text.splitlines():
        heading = re.sub(r'^\s*(?:\d+(?:\.\d+)*[.)]?\s+)?', '', line.strip()).rstrip(':')
        match = script('extract_citations').heading(heading)
        if match:
            current = match[0]
        elif current in sections:
            sections[current] += line + '\n'
        elif current is None:
            front.append(line)
    return {'sections': sections, 'unassigned': '\n'.join(front).strip(),
            'warnings': [] if any(v.strip() for v in sections.values()) else ['헤딩을 찾지 못했습니다. 섹션별 입력으로 직접 나누세요.']}


def section_analysis(sections):
    engine, lit = script('hedge_verbs'), script('litreview')
    result = []
    for key in NAMES:
        text = sections.get(key, '')
        if not text.strip():
            continue
        doc = language()(lit.strip_markers(markers(text))[0])
        words = sum(engine.is_word(t) for t in doc)
        verbs = [t for t in doc if t.pos_ == 'VERB' and t.lemma_.isalpha() and len(t.lemma_) > 1]
        counts = Counter(t.lemma_.lower() for t in verbs)
        sents = list(doc.sents)
        candidates = []
        for si, sent in enumerate(sents):
            for token in sent:
                if token.pos_ == 'VERB' and token.lemma_.lower() in engine.HEDGE_VERB_CANDIDATES:
                    label, reason = engine.prejudge(token)
                    candidates.append(dict(id=f'{key}:{token.idx}', lemma=token.lemma_.lower(), sentence=sent.text,
                                           previous=sents[si-1].text if si else '', next=sents[si+1].text if si+1<len(sents) else '',
                                           automatic=label, label=label, reason=reason, judgment=None, note=''))
        result.append(dict(key=key, section=NAMES[key], words=words, verbs=len(verbs), counts=dict(counts),
                           top=[dict(lemma=l,count=n,rate=round(n/words*10000,1)) for l,n in sorted(counts.items(),key=lambda v:(-v[1],v[0]))[:10]], candidates=candidates))
    return result


def literature(sections):
    """Use the repository's workbook pipeline; retain all editable researcher fields."""
    lit = script('litreview')
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        directory = Path(tmp)
        corpus = directory / 'corpus'
        corpus.mkdir()
        (corpus/'manuscript.json').write_text(json.dumps(dict(id='manuscript',year=None,n_refs=None,sections={k:[markers(p) for p in paragraphs(v)] for k,v in sections.items()})),encoding='utf-8')
        path = directory/'literature.xlsx'
        lit.build(str(corpus),str(path))
        wb=load_workbook(path)
        tables={}
        for key,name in [('intro',lit.INTRO_SHEET),('rd',lit.RD_SHEET)]:
            tables[key]=[list(r) for r in wb[name].iter_rows(min_row=2,values_only=True) if r[0] is not None]
        tables['sizes']=[list(r) for r in wb[lit.SIZE_SHEET].iter_rows(min_row=2,values_only=True) if r[0]]
        return summarize_literature(tables)


def summarize_literature(tables):
    lit=script('litreview')
    topics={}; intro_refs=set(); rd_refs=set(); funcs=Counter(); integral=0
    for row in tables['intro']:
        refs={int(x) for x in str(row[5] or '').split(', ') if x}
        intro_refs |= refs
        integral += row[7]=='Y'
        topic=lit.pick(row[13],row[11],row[10])
        if not topic.startswith('(제외)'):
            topic_key=f'{row[1]}|{topic}'
            entry=topics.setdefault(topic_key,dict(name=topic,paper=row[1],refs=set(),sentences=0,organization=tables.get('organizations',{}).get(topic_key)))
            entry['refs'] |= refs
            entry['sentences'] += 1
    for row in tables['rd']:
        rd_refs |= {int(x) for x in str(row[4] or '').split(', ') if x}
        funcs[lit.pick(row[14],row[12],row[10])] += 1
    sizes=tables.get('sizes',[])
    intro_words=sum(r[3] for r in sizes if r[2]=='서론')
    rd_words=sum(r[3] for r in sizes if r[2]=='결과·논의')
    intro_marks=sum(r[6] for r in sizes if r[2]=='서론')
    rd_marks=sum(r[6] for r in sizes if r[2]=='결과·논의')
    return dict(**tables,topics=[dict(name=t['name'],paper=t['paper'],refs=sorted(t['refs']),sentences=t['sentences'],organization=t['organization']) for t in topics.values()],
                summary=dict(intro_density=intro_marks/intro_words*100 if intro_words else None,rd_density=rd_marks/rd_words*1000 if rd_words else None,intro_marks=intro_marks,rd_marks=rd_marks,topics=len(topics),refs_per_topic=statistics.mean([len(t['refs']) for t in topics.values()]) if topics else 0,
                             intro_refs=len(intro_refs),rd_refs=len(rd_refs),new_refs=len(rd_refs-intro_refs),
                             integral_pct=integral/len(tables['intro'])*100 if tables['intro'] else 0,**{'function:'+f:funcs[f]/len(tables['rd'])*100 if tables['rd'] else None for f in lit.FUNCS},functions=dict(funcs)),
                source='자동 토픽은 문단 단위이며 의미 토픽·조직 유형은 연구자 검토로 확정합니다.')


def distribution(values):
    values=[v for v in values if v is not None]
    return dict(n=len(values),mean=statistics.mean(values) if values else None,min=min(values) if values else None,max=max(values) if values else None)


def compare(sections, baseline):
    for sec in sections:
        hedge=sum((c.get('judgment') or c['automatic'])=='Hedge' for c in sec['candidates'])
        sec['hedge']=dict(count=hedge,per10k=hedge/sec['words']*10000 if sec['words'] else None,per100v=hedge/sec['verbs']*100 if sec['verbs'] else None)
        sec['comparison']=None
        papers=[p['sections'][sec['key']] for p in baseline.get('papers',[]) if sec['key'] in p['sections'] and p['sections'][sec['key']]['words']>0] if baseline else []
        if not papers:
            continue
        lemmas=set(l for p in papers for l in p['counts'])
        reference=[]
        for lemma in lemmas:
            values=[p['counts'].get(lemma,0)/p['words']*10000 for p in papers]
            reference.append(dict(lemma=lemma,papers=sum(p['counts'].get(lemma,0)>0 for p in papers),rate=statistics.mean(values)))
        reference.sort(key=lambda r:(-r['papers'],-r['rate'],r['lemma']))
        top=reference[:10]; top_set={r['lemma'] for r in top}
        sec['comparison']=dict(n=len(papers),top=top,overlap=sum(r['lemma'] in top_set for r in sec['top']),
            coverage=sum(sec['counts'].get(r['lemma'],0)>0 for r in top),missing=[r['lemma'] for r in top if not sec['counts'].get(r['lemma'])],
            rare=[l for l,n in sec['counts'].items() if n>=2 and sum(p['counts'].get(l,0)>0 for p in papers)<=1],
            words=distribution([p['words'] for p in papers]),
            hedge={k:distribution([p['hedge'][k] for p in papers]) for k in ('count','per10k','per100v')},
            verbs={r['lemma']:r for r in reference})
    return sections
