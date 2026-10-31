from concurrent.futures import ProcessPoolExecutor
from contextlib import asynccontextmanager
from datetime import datetime, timezone
import importlib.util
import hashlib
import json
import os
from pathlib import Path
import threading
from uuid import uuid4

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Literal

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
JOBS = DATA / 'jobs'
SCRIPTS = ROOT / 'GNU_AI_Team' / 'hedge-verb-sections' / 'scripts'
LOCK = threading.RLock()
MAX_FILE = 50 * 1024 * 1024
from backend import stage1


def now():
    return datetime.now(timezone.utc).isoformat()


def folder(job_id):
    try:
        from uuid import UUID
        if str(UUID(job_id)) != job_id:
            raise ValueError()
    except ValueError:
        raise HTTPException(404, '작업을 찾을 수 없습니다')
    return JOBS / job_id


def save(path, value):
    temporary = path.with_name(path.name + '.' + str(uuid4()) + '.tmp')
    payload=json.dumps(value,ensure_ascii=False,indent=2).encode('utf-8')
    temporary.write_bytes(payload)
    os.replace(temporary,path)
    if hashlib.sha256(path.read_bytes()).digest()!=hashlib.sha256(payload).digest():
        raise OSError('JSON 저장 후 해시 검증에 실패했습니다')


def read(job_id):
    path = folder(job_id) / 'job.json'
    if not path.exists():
        raise HTTPException(404, '작업을 찾을 수 없습니다')
    return json.loads(path.read_text(encoding='utf-8'))


def load_script(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def analyze(job_id):
    directory = folder(job_id)
    job = read(job_id)
    try:
        job.update(status='running', updated_at=now())
        save(directory / 'job.json', job)
        if job.get('mode')=='stage1':
            try:
                value=stage1.generate(directory,job['baseline_name'],str(uuid4()),now())
                value['input_files']=job['files']
                destination=DATA/'baselines';destination.mkdir(parents=True,exist_ok=True)
                path=destination/(value['id']+'.json');save(path,value)
                assert json.loads(path.read_text(encoding='utf-8'))==value
                job.update(status='completed',baseline_id=value['id'],updated_at=now())
            finally:
                gate=directory/'stage1'/'extract_gate.json'
                if gate.exists():
                    job['gate']=json.loads(gate.read_text(encoding='utf-8'))
            save(directory/'job.json',job)
            return
        extractor = engine.script('extract_citations')
        corpus = directory / 'corpus'
        corpus.mkdir(exist_ok=True)
        for pdf in sorted((directory / 'pdf').glob('*.pdf')):
            document = extractor.process(str(pdf))
            document['sections'] = {k:[engine.markers(p) for p in v] for k,v in document['sections'].items()}
            if not document['sections']:
                raise ValueError(f'{pdf.name}: 섹션을 인식하지 못했습니다. 텍스트 PDF와 헤딩을 확인하세요.')
            save(corpus / f'{pdf.stem}.json', dict(document, sections={k:[engine.script('litreview').strip_markers(p)[0] for p in v] for k,v in document['sections'].items()}))
            citation_corpus = directory / 'citation_corpus'
            citation_corpus.mkdir(exist_ok=True)
            save(citation_corpus / f'{pdf.stem}.json', document)
        verbs = load_script('hedge_verbs')
        verbs.build(str(corpus), str(directory / 'result.xlsx'))
        from openpyxl import load_workbook
        workbook = load_workbook(directory / 'result.xlsx')
        keys = ['id', 'paper', 'section', 'sentence_number', 'form', 'lemma', 'previous', 'sentence', 'next', 'automatic', 'reason', 'review', 'review_note', 'judgment', 'note']
        candidates = [dict(zip(keys, row)) for row in workbook['판정작업지'].iter_rows(min_row=2, values_only=True) if row[0] is not None]
        save(directory / 'candidates.json', candidates)
        lit = engine.script('litreview')
        lit.build(str(directory/'citation_corpus'), str(directory/'literature.xlsx'))
        lit.baseline(str(directory/'literature.xlsx'),str(directory/'baseline_lit.json'))
        lit_wb = load_workbook(directory/'literature.xlsx')
        tables = {key:[list(r) for r in lit_wb[sheet].iter_rows(min_row=2,values_only=True) if r[0] is not None]
                  for key,sheet in [('intro',lit.INTRO_SHEET),('rd',lit.RD_SHEET)]}
        tables['sizes']=[list(r) for r in lit_wb[lit.SIZE_SHEET].iter_rows(min_row=2,values_only=True) if r[0]]
        save(directory/'literature.json',engine.summarize_literature(tables))
        job['extraction'] = [dict(file=d['id'], sections=list(d['sections']), missing=[k for k in engine.NAMES if not d['sections'].get(k)], n_refs=d.get('n_refs'))
                             for d in [json.loads(p.read_text(encoding='utf-8')) for p in (directory/'citation_corpus').glob('*.json')]]
        job.update(status='completed', candidate_count=len(candidates), updated_at=now())
    except Exception as exc:
        job.update(status='failed', error=str(exc), updated_at=now())
    save(directory / 'job.json', job)


@asynccontextmanager
async def lifespan(app):
    JOBS.mkdir(parents=True, exist_ok=True)
    baseline_directory=DATA/'baselines';baseline_directory.mkdir(parents=True,exist_ok=True)
    builtin=baseline_directory/(stage1.BUILTIN_ID+'.json')
    value=stage1.import_baseline(stage1.EXAMPLE,'무기화학 · IC 2010–2015 · Stage 1',stage1.BUILTIN_ID,'2026-10-31T00:00:00+00:00',gate=json.loads((stage1.EXAMPLE/'extract_gate.json').read_text(encoding='utf-8')))
    if not builtin.exists() or json.loads(builtin.read_text(encoding='utf-8'))!=value:
        save(builtin,value)
    if json.loads(builtin.read_text(encoding='utf-8'))!=value:
        raise RuntimeError('포함 기준선의 저장 내용이 원본과 다릅니다')
    for bundled in (ROOT/'baselines_markdown').glob('*/baseline.json'):
        record=json.loads(bundled.read_text(encoding='utf-8'))
        target=baseline_directory/(folder(record['id']).name+'.json')
        if not target.exists():
            save(target,record)
    for path in JOBS.glob('*/job.json'):
        job = json.loads(path.read_text(encoding='utf-8'))
        if job['status'] in ('queued', 'running'):
            job.update(status='failed', error='서버가 종료되어 분석이 중단되었습니다', updated_at=now())
            save(path, job)
    app.state.pool = ProcessPoolExecutor(max_workers=1)
    yield
    app.state.pool.shutdown(wait=True)


app = FastAPI(title='Hedge Verb Analysis', lifespan=lifespan)


@app.get('/api/health')
def health():
    return {'status': 'ok'}


@app.post('/api/jobs', status_code=202)
async def create_job(files: list[UploadFile] = File(...), mode: Literal['legacy','stage1'] = Form('legacy'), baseline_name: str = Form('새 Stage 1 기준선')):
    if mode=='stage1' and not 2 <= len(files) <= 10:
        raise HTTPException(400,'Stage 1 기준선은 PDF 2~10편으로 생성합니다')
    if not 1 <= len(files) <= 20:
        raise HTTPException(400, 'PDF는 1~20개까지 업로드하세요')
    job_id = str(uuid4())
    directory = folder(job_id)
    (directory / 'pdf').mkdir(parents=True)
    names = []
    try:
        for index, upload in enumerate(files, 1):
            if not (upload.filename or '').lower().endswith('.pdf'):
                raise HTTPException(400, 'PDF 파일만 지원합니다')
            target = directory / 'pdf' / f'{index:03d}.pdf'
            size = 0
            with target.open('wb') as output:
                while chunk := await upload.read(1024 * 1024):
                    if size == 0 and not chunk.startswith(b'%PDF-'):
                        raise HTTPException(400, '유효한 PDF 헤더가 없습니다')
                    size += len(chunk)
                    if size > MAX_FILE:
                        raise HTTPException(413, '파일당 최대 50MB입니다')
                    output.write(chunk)
            if not size:
                raise HTTPException(400, '빈 파일입니다')
            names.append({'file': target.name, 'original_name': upload.filename, 'sha256': hashlib.sha256(target.read_bytes()).hexdigest()})
        if mode=='stage1' and len({f['sha256'] for f in names})!=len(names):
            raise HTTPException(400,'동일한 PDF를 중복 업로드할 수 없습니다')
        job = dict(id=job_id, mode=mode,baseline_name=baseline_name.strip() or '새 Stage 1 기준선', status='queued', created_at=now(), updated_at=now(), files=names)
        save(directory / 'job.json', job)
        app.state.pool.submit(analyze, job_id)
        return job
    except Exception:
        import shutil
        shutil.rmtree(directory)
        raise
    finally:
        for upload in files:
            await upload.close()


@app.get('/api/jobs')
def list_jobs():
    with LOCK:
        return sorted([json.loads(p.read_text(encoding='utf-8')) for p in JOBS.glob('*/job.json')], key=lambda j: j['created_at'], reverse=True)


@app.get('/api/jobs/{job_id}')
def get_job(job_id: str):
    return read(job_id)


def completed(job_id):
    if read(job_id)['status'] != 'completed':
        raise HTTPException(409, '분석 완료 후 이용할 수 있습니다')
    return folder(job_id)


@app.get('/api/jobs/{job_id}/candidates')
def candidates(job_id: str):
    with LOCK:
        directory = completed(job_id)
        return json.loads((directory / 'candidates.json').read_text(encoding='utf-8'))


class Judgment(BaseModel):
    judgment: Literal['Hedge', 'N', 'U'] | None = None
    note: str = ''


@app.patch('/api/jobs/{job_id}/candidates/{candidate_id}')
def update_judgment(job_id: str, candidate_id: int, body: Judgment):
    with LOCK:
        directory = completed(job_id)
        rows = json.loads((directory / 'candidates.json').read_text(encoding='utf-8'))
        row = next((r for r in rows if r['id'] == candidate_id), None)
        if row is None:
            raise HTTPException(404, '후보를 찾을 수 없습니다')
        row.update(judgment=body.judgment, note=body.note)
        save(directory / 'candidates.json', rows)
        return row


@app.get('/api/jobs/{job_id}/download')
def download(job_id: str):
    with LOCK:
        directory = completed(job_id)
        from openpyxl import load_workbook
        path = directory / 'result.xlsx'
        workbook = load_workbook(path)
        rows = json.loads((directory / 'candidates.json').read_text(encoding='utf-8'))
        for row in rows:
            workbook['판정작업지'].cell(row['id'] + 1, 14, row['judgment'] or '')
            workbook['판정작업지'].cell(row['id'] + 1, 15, row['note'])
        # Each download has its own snapshot, so concurrent responses cannot overwrite it.
        snapshot = directory / f'export-{uuid4()}.xlsx'
        workbook.save(snapshot)
        load_script('hedge_verbs').stats(str(snapshot))
        from starlette.background import BackgroundTask
        return FileResponse(snapshot, filename=f'hedge-{job_id}.xlsx', background=BackgroundTask(snapshot.unlink, missing_ok=True))


from backend import analysis as engine
from fastapi import Form


class Manuscript(BaseModel):
    sections: dict[Literal['INTRO', 'M', 'RD', 'CONCL'], str]
    baseline_id: str | None = None


def baseline_read(baseline_id):
    if not baseline_id:
        return None
    path = DATA / 'baselines' / (folder(baseline_id).name + '.json')
    if not path.exists():
        raise HTTPException(404, '분야 기준선을 찾을 수 없습니다')
    return json.loads(path.read_text(encoding='utf-8'))


@app.post('/api/manuscript')
def manuscript(body: Manuscript):
    if sum(map(len, body.sections.values())) > 200000:
        raise HTTPException(413, '원고는 200,000자 이내로 입력하세요')
    if not any(v.strip() for v in body.sections.values()):
        raise HTTPException(400, '원고를 입력하세요')
    baseline = baseline_read(body.baseline_id)
    try:
        stage_result=stage1.manuscript(body.sections,baseline) if baseline and baseline.get('kind')=='stage1' else None
    except (ValueError,TimeoutError) as exc:
        raise HTTPException(422,str(exc))
    result = dict(stage1=stage_result,id=str(uuid4()), created_at=now(), baseline_id=body.baseline_id,
                  baseline_name=baseline['name'] if baseline else None, sections=engine.section_analysis(body.sections),
                  literature=engine.literature(body.sections), input=body.sections)
    engine.compare(result['sections'], baseline if not stage_result else None)
    result['baseline_literature'] = {k:engine.distribution([p['literature'].get(k) for p in baseline['papers']]) for k in result['literature']['summary'] if k!='functions'} if baseline and not stage_result else None
    directory = DATA / 'manuscripts'
    directory.mkdir(parents=True, exist_ok=True)
    with LOCK:
        save(directory / (result['id'] + '.json'), result)
    return result


class SplitInput(BaseModel):
    text: str


@app.post('/api/split')
def split_text(body: SplitInput):
    if len(body.text)>200000:
        raise HTTPException(413, '원고는 200,000자 이내로 입력하세요')
    return engine.split_sections(body.text)


@app.post('/api/import')
async def import_document(file: UploadFile = File(...)):
    import io, zipfile
    from xml.etree import ElementTree
    try:
        content = await file.read(MAX_FILE + 1)
        if len(content)>MAX_FILE:
            raise HTTPException(413, '파일당 최대 50MB입니다')
        extension = Path(file.filename or '').suffix.lower()
        if extension == '.txt':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('cp949')
        elif extension == '.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if archive.getinfo('word/document.xml').file_size>10*1024*1024:
                    raise ValueError('DOCX 본문이 너무 큽니다')
                root = ElementTree.fromstring(archive.read('word/document.xml'))
                ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                text = '\n\n'.join(''.join(t.text or '' for t in para.findall('.//w:t',ns)) for para in root.findall('.//w:p',ns))
        elif extension == '.pdf':
            import pymupdf
            with pymupdf.open(stream=content,filetype='pdf') as doc:
                text = '\n\n'.join(page.get_text(sort=True) for page in doc)
            if not text.strip():
                raise ValueError('텍스트를 추출할 수 없습니다. 스캔 PDF는 OCR이 필요합니다')
        else:
            raise HTTPException(400, 'TXT, DOCX, PDF만 지원합니다')
        if len(text)>200000:
            raise HTTPException(413, '추출된 원고가 200,000자를 초과합니다')
        return dict(text=text, **engine.split_sections(text))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f'파일을 읽지 못했습니다: {exc}')
    finally:
        await file.close()


@app.get('/api/baselines')
def baselines():
    return [dict(id=b['id'],name=b['name'],kind=b.get('kind','legacy'),group=b.get('group'),input_kind=b.get('input_kind'),needs_review=b.get('needs_review',True),papers=len(b['papers']),created_at=b['created_at'],source_jobs=b['source_jobs']) for b in
            [json.loads(p.read_text(encoding='utf-8')) for p in (DATA/'baselines').glob('*.json')]]


class BaselineInput(BaseModel):
    name: str
    job_ids: list[str]


@app.post('/api/baselines', status_code=201)
def create_baseline(body: BaselineInput):
    if not body.name.strip() or not body.job_ids:
        raise HTTPException(400, '분야명과 완료된 PDF 분석 작업을 선택하세요')
    papers=[]; seen=set(); automatic=0
    with LOCK:
        for jid in dict.fromkeys(body.job_ids):
            directory=completed(jid)
            if read(jid).get('mode')=='stage1':
                raise HTTPException(400,'Stage 1 작업은 완료 시 자체 기준선을 생성합니다')
            rows=json.loads((directory/'candidates.json').read_text(encoding='utf-8'))
            docs=[json.loads(p.read_text(encoding='utf-8')) for p in (directory/'corpus').glob('*.json')]
            docs.sort(key=lambda d:(d['year'] or 0,d['id']))
            lit=job_literature(directory)
            for index,doc in enumerate(docs,1):
                pid=f'P{index:02d}'
                source=next((f for f in read(jid)['files'] if Path(f['file']).stem==doc['id']),None)
                digest=source.get('sha256') if source else None
                digest=digest or hashlib.sha256((directory/'pdf'/(doc['id']+'.pdf')).read_bytes()).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                automatic+=sum(not (r.get('judgment') or r.get('review')) for r in rows if r['paper']==pid)
                sections=engine.section_analysis({k:'\n\n'.join(v) for k,v in doc['sections'].items()})
                for sec in sections:
                    original=[r for r in rows if r['paper']==pid and r['section']==sec['section']]
                    # Corpus and workbook share the same extraction and token order.
                    if len(sec['candidates'])!=len(original):
                        raise HTTPException(409,'후보 수가 달라 기준선을 만들 수 없습니다')
                    for candidate,review in zip(sec['candidates'],original):
                        if candidate['lemma']!=review['lemma']:
                            raise HTTPException(409,'후보 순서가 달라 기준선을 만들 수 없습니다')
                        candidate['judgment']=review.get('judgment') or review.get('review')
                engine.compare(sections,None)
                table={k:[r for r in lit[k] if r[1]==pid] for k in ('intro','rd')}
                table['organizations']=lit.get('organizations',{})
                table['sizes']=[r for r in lit.get('sizes',[]) if r[0]==pid]
                summary=engine.summarize_literature(table)['summary']
                if not doc['sections'].get('INTRO'):
                    for key in ('topics','refs_per_topic','intro_refs','intro_marks','intro_density','integral_pct'):
                        summary[key]=None
                if not doc['sections'].get('RD'):
                    for key in ('rd_refs','rd_marks','rd_density','new_refs'):
                        summary[key]=None
                    for function in engine.script('litreview').FUNCS:
                        summary['function:'+function]=None
                if not doc['sections'].get('INTRO'):
                    summary['new_refs']=None
                papers.append(dict(id=jid+':'+doc['id'],sha256=digest,sections={s['key']:s for s in sections},literature=summary))
        result=dict(id=str(uuid4()),name=body.name.strip(),created_at=now(),source_jobs=list(dict.fromkeys(body.job_ids)),papers=papers,unreviewed_hedge_candidates=automatic,judgment_source='연구자 > 검토 > 자동; 선행연구 토픽·기능은 수동 검토 전 예비판정 포함')
        directory=DATA/'baselines';directory.mkdir(exist_ok=True,parents=True)
        save(directory/(result['id']+'.json'),result)
    return dict(id=result['id'],name=result['name'],papers=len(papers),unreviewed_hedge_candidates=automatic)


class LiteratureReview(BaseModel):
    area: Literal['intro','rd','organization']
    row_id: int | None = None
    value: str | None = None
    topic: str | None = None
    note: str = ''


def job_literature(directory):
    path=directory/'literature.json'
    if not path.exists():
        raise HTTPException(409,'이 작업에는 선행연구 결과가 없습니다. PDF를 다시 분석하세요')
    return json.loads(path.read_text(encoding='utf-8'))


def edit_literature(tables,body):
    if body.area=='organization':
        if body.value not in ('주제 통합','논문별 나열','혼합',None):
            raise HTTPException(400,'조직 유형을 확인하세요')
        valid_topics={t['paper']+'|'+t['name'] for t in tables['topics']}
        if body.topic not in valid_topics:
            raise HTTPException(404,'토픽을 찾을 수 없습니다')
        tables.setdefault('organizations',{})[body.topic]=body.value
    else:
        row=next((r for r in tables[body.area] if r[0]==body.row_id),None)
        if row is None:
            raise HTTPException(404,'인용 문장을 찾을 수 없습니다')
        if body.area=='rd' and body.value not in engine.script('litreview').FUNCS+[None]:
            raise HTTPException(400,'인용 기능을 확인하세요')
        column=13 if body.area=='intro' else 14
        row[column]=body.value
        row[column+1]=body.note
    return engine.summarize_literature({k:tables[k] for k in ('intro','rd','organizations','sizes') if k in tables})


@app.get('/api/jobs/{job_id}/literature')
def get_literature(job_id:str):
    with LOCK:
        return job_literature(completed(job_id))


@app.patch('/api/jobs/{job_id}/literature')
def patch_literature(job_id:str,body:LiteratureReview):
    with LOCK:
        directory=completed(job_id)
        value=edit_literature(job_literature(directory),body)
        save(directory/'literature.json',value)
        return value


@app.patch('/api/manuscripts/{manuscript_id}/literature')
def patch_manuscript_literature(manuscript_id:str,body:LiteratureReview):
    path=DATA/'manuscripts'/(folder(manuscript_id).name+'.json')
    with LOCK:
        if not path.exists():
            raise HTTPException(404,'원고 분석을 찾을 수 없습니다')
        result=json.loads(path.read_text(encoding='utf-8'))
        result['literature']=edit_literature(result['literature'],body)
        save(path,result)
        return result['literature']


@app.get('/api/jobs/{job_id}/literature/download')
def download_literature(job_id:str):
    from openpyxl import load_workbook
    from starlette.background import BackgroundTask
    with LOCK:
        directory=completed(job_id);tables=job_literature(directory)
        wb=load_workbook(directory/'literature.xlsx')
        lit=engine.script('litreview')
        for key,sheet,col in [('intro',lit.INTRO_SHEET,14),('rd',lit.RD_SHEET,15)]:
            for row in tables[key]:
                wb[sheet].cell(row[0]+1,col,row[col-1] or '')
                wb[sheet].cell(row[0]+1,col+1,row[col] or '')
        for row in wb[lit.TOPIC_SHEET].iter_rows(min_row=2):
            topic=row[1].value
            topic_key=f'{row[0].value}|{topic}'
            if topic_key in tables.get('organizations',{}):
                row[6].value=tables['organizations'][topic_key]
        path=directory/f'lit-export-{uuid4()}.xlsx';wb.save(path);lit.stats(str(path))
        if tables.get('organizations'):
            wb=load_workbook(path)
            for row in wb[lit.TOPIC_SHEET].iter_rows(min_row=2):
                key=f'{row[0].value}|{row[1].value}'
                if key in tables['organizations']:
                    row[6].value=tables['organizations'][key]
            wb.save(path)
            lit.stats(str(path))
        return FileResponse(path,filename='literature.xlsx',background=BackgroundTask(path.unlink,missing_ok=True))


@app.get('/')
def index():
    return FileResponse(ROOT / 'frontend' / 'index.html')


@app.get('/app.js')
def javascript():
    return FileResponse(ROOT/'frontend'/'app.js',media_type='application/javascript')


class ManuscriptJudgment(Judgment):
    candidate_id: str


@app.patch('/api/manuscripts/{manuscript_id}/candidates')
def patch_manuscript_candidate(manuscript_id:str,body:ManuscriptJudgment):
    path=DATA/'manuscripts'/(folder(manuscript_id).name+'.json')
    with LOCK:
        if not path.exists():
            raise HTTPException(404,'원고 분석을 찾을 수 없습니다')
        result=json.loads(path.read_text(encoding='utf-8'))
        candidate=next((c for s in result['sections'] for c in s['candidates'] if c['id']==body.candidate_id),None)
        if candidate is None:
            raise HTTPException(404,'후보를 찾을 수 없습니다')
        candidate.update(judgment=body.judgment,note=body.note)
        baseline=baseline_read(result['baseline_id'])
        if baseline and baseline.get('kind')=='stage1':
            raise HTTPException(409,'Stage 1 규칙 결과는 기존 어휘동사 판정으로 수정할 수 없습니다')
        engine.compare(result['sections'],baseline)
        save(path,result)
        return result


@app.get('/api/baselines/{baseline_id}')
def get_baseline(baseline_id:str):
    return baseline_read(baseline_id)


@app.get('/api/jobs/{job_id}/stage1/download')
def stage1_artifacts(job_id:str):
    import zipfile
    from starlette.background import BackgroundTask
    job=read(job_id)
    if job['status'] not in ('completed','failed'):
        raise HTTPException(409,'분석 완료 또는 실패 후 결과를 내려받으세요')
    directory=folder(job_id)/'stage1'
    if not directory.exists():
        raise HTTPException(404,'Stage 1 결과가 없습니다')
    snapshot=folder(job_id)/f'stage1-{uuid4()}.zip'
    with zipfile.ZipFile(snapshot,'w',zipfile.ZIP_DEFLATED) as archive:
        for path in directory.rglob('*'):
            if path.is_file():
                archive.write(path,path.relative_to(directory))
    return FileResponse(snapshot,filename='stage1-baseline.zip',background=BackgroundTask(snapshot.unlink,missing_ok=True))
