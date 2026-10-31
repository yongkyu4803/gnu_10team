import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
import zipfile

import pymupdf
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend import main


class Workflow(unittest.TestCase):
    def test_analysis_review_baseline_and_download(self):
        old_data,old_jobs=main.DATA,main.JOBS
        with tempfile.TemporaryDirectory() as tmp:
            main.DATA=Path(tmp);main.JOBS=Path(tmp)/'jobs'
            try:
                with TestClient(main.app) as client:
                    main.app.state.pool.shutdown()
                    main.app.state.pool=ThreadPoolExecutor(max_workers=1)
                    self.assertEqual(client.get('/').status_code,200)
                    self.assertEqual(client.get('/app.js').status_code,200)
                    self.assertEqual(client.post('/api/manuscript',json={'sections':{}}).status_code,400)
                    sections={'INTRO':'Earlier studies suggest that the solvent controls the reaction [1–3].\n\nSmith et al. reported a new pathway [4].','RD':'Our results are consistent with previous observations [1]. The reaction follows a published procedure [5].','M':'The mixture was heated and stirred.','CONCL':'These results suggest a new pathway.'}
                    result=client.post('/api/manuscript',json={'sections':sections}).json()
                    self.assertEqual(result['literature']['summary']['topics'],2)
                    self.assertEqual(result['literature']['summary']['intro_refs'],4)
                    self.assertEqual(result['literature']['summary']['new_refs'],1)
                    mid=result['id']
                    for row in result['literature']['intro']:
                        response=client.patch(f'/api/manuscripts/{mid}/literature',json={'area':'intro','row_id':row[0],'value':'선행 반응 경로'})
                        self.assertEqual(response.status_code,200,response.text)
                    self.assertEqual(response.json()['summary']['topics'],1)
                    response=client.patch(f'/api/manuscripts/{mid}/literature',json={'area':'organization','topic':'P01|선행 반응 경로','value':'주제 통합'})
                    self.assertEqual(response.status_code,200,response.text)
                    self.assertEqual(response.json()['topics'][0]['organization'],'주제 통합')
                    empty=client.post('/api/manuscript',json={'sections':{'INTRO':'Water is a liquid.'}})
                    self.assertEqual(empty.status_code,200,empty.text)
                    self.assertEqual(empty.json()['literature']['summary']['topics'],0)
                    self.assertEqual(empty.json()['sections'][0]['candidates'],[])
                    row=result['literature']['rd'][0]
                    response=client.patch(f'/api/manuscripts/{mid}/literature',json={'area':'rd','row_id':row[0],'value':'해석 근거'})
                    self.assertEqual(response.json()['summary']['functions']['해석 근거'],1)
                    candidate=result['sections'][0]['candidates'][0]
                    before=result['sections'][0]['hedge']['count']
                    response=client.patch(f'/api/manuscripts/{mid}/candidates',json={'candidate_id':candidate['id'],'judgment':'N'})
                    self.assertEqual(response.json()['sections'][0]['hedge']['count'],before-1)
                    split=client.post('/api/split',json={'text':'Title\n1. INTRODUCTION\nIntro text\n2. RESULTS AND DISCUSSION\nResults text\nREFERENCES\n(1) A paper'}).json()
                    self.assertIn('Intro text',split['sections']['INTRO'])
                    self.assertNotIn('A paper',split['sections']['RD'])
                    for extension,content in [('txt',b'INTRODUCTION\nExample input text.')]:
                        response=client.post('/api/import',files={'file':('draft.'+extension,content)})
                        self.assertEqual(response.status_code,200,response.text)
                    docx=io.BytesIO()
                    with zipfile.ZipFile(docx,'w') as archive:
                        archive.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Example manuscript.</w:t></w:r></w:p></w:body></w:document>')
                    self.assertIn('Example manuscript',client.post('/api/import',files={'file':('draft.docx',docx.getvalue())}).json()['text'])
                    doc=pymupdf.open();page=doc.new_page();y=60
                    for heading,text in [('INTRODUCTION',sections['INTRO']),('EXPERIMENTAL SECTION',sections['M']),('RESULTS AND DISCUSSION',sections['RD']),('CONCLUSION',sections['CONCL']),('REFERENCES','(1) Earlier observations in chemistry and reaction pathways.\n(2) Another reference concerning solvent and reactive intermediates.\n(3) A study of molecular reactivity and hydrogen bonds.\n(4) A new pathway in molecular chemistry and synthesis.\n(5) Procedures for organic synthesis and experiments.')]:
                        page.insert_text((50,y),heading,fontsize=12);y+=22
                        rect=pymupdf.Rect(50,y,540,y+85)
                        page.insert_textbox(rect,text,fontsize=10)
                        y+=105
                    content=doc.tobytes();doc.close()
                    self.assertEqual(client.post('/api/import',files={'file':('paper.pdf',content)}).status_code,200)
                    response=client.post('/api/jobs',files={'files':('paper.pdf',content,'application/pdf')})
                    self.assertEqual(response.status_code,202,response.text)
                    jid=response.json()['id']
                    for _ in range(150):
                        job=client.get('/api/jobs/'+jid).json()
                        if job['status'] in ('completed','failed'):break
                        time.sleep(.1)
                    self.assertEqual(job['status'],'completed',job)
                    rows=client.get(f'/api/jobs/{jid}/candidates').json()
                    self.assertTrue(rows)
                    response=client.patch(f'/api/jobs/{jid}/candidates/{rows[0]["id"]}',json={'judgment':'N','note':'검토 완료'})
                    self.assertEqual(response.status_code,200)
                    response=client.get(f'/api/jobs/{jid}/download')
                    wb=load_workbook(io.BytesIO(response.content))
                    self.assertEqual(wb['판정작업지'].cell(2,14).value,'N')
                    lit=client.get(f'/api/jobs/{jid}/literature').json()
                    self.assertTrue(lit['intro'])
                    response=client.patch(f'/api/jobs/{jid}/literature',json={'area':'intro','row_id':lit['intro'][0][0],'value':'반응 경로'})
                    self.assertEqual(response.status_code,200,response.text)
                    response=client.patch(f'/api/jobs/{jid}/literature',json={'area':'organization','topic':'P01|반응 경로','value':'주제 통합'})
                    self.assertEqual(response.status_code,200,response.text)
                    response=client.get(f'/api/jobs/{jid}/literature/download')
                    wb=load_workbook(io.BytesIO(response.content))
                    self.assertEqual(wb['서론_인용문장'].cell(2,14).value,'반응 경로')
                    self.assertEqual(wb['서론_토픽'].cell(2,7).value,'주제 통합')
                    response=client.post('/api/baselines',json={'name':'검증용 코퍼스','job_ids':[jid,jid]})
                    self.assertEqual(response.status_code,201,response.text)
                    baseline=response.json();self.assertEqual(baseline['papers'],1)
                    self.assertIn('unreviewed_hedge_candidates',baseline)
                    response=client.post('/api/manuscript',json={'sections':sections,'baseline_id':baseline['id']})
                    self.assertEqual(response.status_code,200,response.text)
                    comparison=response.json()['sections'][0]['comparison']
                    self.assertEqual(comparison['n'],1)
                    self.assertGreater(comparison['coverage'],0)
                    self.assertIsNotNone(response.json()['baseline_literature'])
                    stored=json.loads((main.DATA/'baselines'/(baseline['id']+'.json')).read_text())
                    self.assertEqual(stored['source_jobs'],[jid])
                    self.assertEqual(client.get('/api/jobs/invalid-id').status_code,404)
                    self.assertEqual(client.post('/api/jobs',files={'files':('x.pdf',b'bad')}).status_code,400)
            finally:
                main.DATA,main.JOBS=old_data,old_jobs


if __name__=='__main__':
    unittest.main()
